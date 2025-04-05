import os
import csv
import shutil
import tempfile
from flask import Flask, render_template, request, redirect, url_for, flash
from werkzeug.utils import secure_filename
from PIL import Image, ImageDraw
from docx import Document
import openpyxl
from pillow_heif import register_heif_opener
from dotenv import load_dotenv
import boto3
from langchain.chat_models import ChatOpenAI
from langchain.prompts import ChatPromptTemplate

# Load environment variables
load_dotenv()

# Flask setup
UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'txt', 'csv', 'docx', 'xlsx', 'heic', 'jpg', 'jpeg', 'png'}
app = Flask(__name__)
app.secret_key = 'supersecretkey'
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# AWS + OpenAI setup
AWS_REGION = os.getenv("AWS_REGION")
AWS_ACCESS_KEY = os.getenv("AWS_ACCESS_KEY")
AWS_SECRET_KEY = os.getenv("AWS_SECRET_KEY")

textract = boto3.client(
    "textract",
    region_name=AWS_REGION,
    aws_access_key_id=AWS_ACCESS_KEY,
    aws_secret_access_key=AWS_SECRET_KEY
)

llm = ChatOpenAI(
    openai_api_key=os.getenv("OPENAI_API_KEY"),
    model_name="gpt-4o-mini"
)

RUBRIC = """
--Data table including mass of ball, pendulum, initial velocity from photo gate, and 5 trials indicating angle of deflection (10 pts)
--Calculation and comparison of velocity from photogate vs. velocity from angular deflection (4pts)
--Calculation of percent difference in velocities (2pts)
--Brief discussion and analysis of results (individual) (4pts)
"""

prompt = ChatPromptTemplate.from_messages([
    ("system", "You are a homework evaluation assistant. Evaluate the extracted text based on the rubric provided. Detect whether the homework contains data tables by identifying chart titles or tabular numeric patterns. Grade based on completion, not correctness. Don't deduct points for errors. If they attempted it, give full credit."),
    ("user", "{input}")
])

# File helpers
def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def convert_docx_to_jpg(docx_path, output_path):
    doc = Document(docx_path)
    img = Image.new("RGB", (1000, 1400), color="white")
    draw = ImageDraw.Draw(img)
    y = 20
    for para in doc.paragraphs:
        draw.text((20, y), para.text, fill="black")
        y += 30
    img.save(output_path)

def convert_xlsx_to_jpg(xlsx_path, output_path):
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    img = Image.new("RGB", (1000, 1400), color="white")
    draw = ImageDraw.Draw(img)
    y = 20
    for row in ws.iter_rows(values_only=True):
        row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
        draw.text((20, y), row_text, fill="black")
        y += 30
    img.save(output_path)

def convert_heic_to_jpg(heic_path, output_path):
    register_heif_opener()
    img = Image.open(heic_path)
    img.save(output_path, "JPEG")

def convert_to_jpg(file_path):
    file_ext = os.path.splitext(file_path)[1].lower()
    jpg_output_path = tempfile.mktemp(suffix=".jpg")
    try:
        if file_ext == ".docx":
            convert_docx_to_jpg(file_path, jpg_output_path)
        elif file_ext == ".xlsx":
            convert_xlsx_to_jpg(file_path, jpg_output_path)
        elif file_ext == ".heic":
            convert_heic_to_jpg(file_path, jpg_output_path)
        else:
            return None
        return jpg_output_path
    except Exception as e:
        print(f"Error converting {file_path}: {e}")
        return None

def extract_text_textract(image_path):
    with open(image_path, "rb") as image_file:
        image_bytes = image_file.read()
    response = textract.detect_document_text(Document={"Bytes": image_bytes})
    return "\n".join([item["Text"] for item in response["Blocks"] if item["BlockType"] == "LINE"])

def detect_graphs_and_tables(text):
    graph_keywords = [" vs ", "graph", "chart", "plot", "trend"]
    table_patterns = ["\t", ",", " | ", " |", "| ", "row", "column", "data"]
    graph_count = sum(1 for keyword in graph_keywords if keyword.lower() in text.lower())
    table_count = sum(1 for line in text.split("\n") if any(p in line for p in table_patterns))
    return graph_count >= 2, table_count >= 2

def evaluate_homework(file_path):
    file_ext = os.path.splitext(file_path)[1].lower()
    if file_ext not in [".jpg", ".jpeg", ".png"]:
        jpg_path = convert_to_jpg(file_path)
        if not jpg_path:
            print(f"Conversion failed for: {file_path}")
            return 0
        image_path = jpg_path
    else:
        image_path = file_path
    extracted_text = extract_text_textract(image_path)
    if not extracted_text:
        return 0
    graphs, tables = detect_graphs_and_tables(extracted_text)
    input_text = f"""Evaluate this homework based on the rubric:\n{RUBRIC}\n\nExtracted Text:\n{extracted_text}\n\nGraphs detected: {'Yes' if graphs else 'No'}\nTables detected: {'Yes' if tables else 'No'}"""
    chain = prompt | llm
    result = chain.invoke({"input": input_text})
    print(f"\nEvaluation Result:\n{result.content}")
    return 20  # Placeholder: adjust this to parse real score from result.content if needed

# Flask routes
@app.route('/', methods=['GET', 'POST'])
def index():
    if request.method == 'POST':
        activity_file = request.files['activity']
        gradebook_file = request.files['gradebook']
        submission_files = request.files.getlist('submissions')

        if os.path.exists(UPLOAD_FOLDER):
            shutil.rmtree(UPLOAD_FOLDER)
        os.makedirs(UPLOAD_FOLDER, exist_ok=True)

        if activity_file and allowed_file(activity_file.filename):
            activity_path = os.path.join(UPLOAD_FOLDER, secure_filename(activity_file.filename))
            activity_file.save(activity_path)
        else:
            flash("Invalid or missing activity file.")
            return redirect(request.url)

        if gradebook_file and allowed_file(gradebook_file.filename):
            gradebook_path = os.path.join(UPLOAD_FOLDER, secure_filename(gradebook_file.filename))
            gradebook_file.save(gradebook_path)
        else:
            flash("Invalid or missing gradebook file.")
            return redirect(request.url)

        submissions_dir = os.path.join(UPLOAD_FOLDER, 'submissions')
        os.makedirs(submissions_dir, exist_ok=True)
        for f in submission_files:
            if f and allowed_file(f.filename):
                f.save(os.path.join(submissions_dir, secure_filename(f.filename)))

        process_grades(gradebook_path, submissions_dir)
        flash("Processing complete and grades updated.")
        return redirect(url_for('index'))

    return render_template('index.html')

def process_grades(gradebook_path, submissions_dir):
    names = []
    with open(gradebook_path, newline='') as csvfile:
        reader = csv.reader(csvfile)
        for row in reader:
            names.append(row[0])  # Assumes names in column A

    updates = []
    for name in names:
        matched_file = next((f for f in os.listdir(submissions_dir) if name.lower() in f.lower()), None)
        if matched_file:
            file_path = os.path.join(submissions_dir, matched_file)
            print(f"Evaluating: {file_path}")
            score = evaluate_homework(file_path)
            updates.append([score])
        else:
            updates.append([0])

    print("Final grades:", updates)

if __name__ == '__main__':
    app.run(debug=True)
